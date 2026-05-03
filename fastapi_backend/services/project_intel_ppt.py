"""Build a PowerPoint deck from a Project Intel AnalysisReport using python-pptx."""

from __future__ import annotations

import re
from io import BytesIO
from typing import Iterable, List

from pptx import Presentation
from pptx.util import Pt

from fastapi_backend.schemas.analysis import AnalysisReport


def _truncate(text: str, max_len: int = 500) -> str:
    text = (text or "").strip()
    if len(text) <= max_len:
        return text
    return text[: max_len - 1] + "…"


def _safe_filename(title: str) -> str:
    base = re.sub(r"[^\w\-]+", "_", (title or "report").strip())[:60] or "report"
    return f"{base}.pptx"


def _add_title_content_slide(prs: Presentation, title: str, body_lines: List[str]) -> None:
    layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = _truncate(title, 120)
    body = slide.placeholders[1]
    tf = body.text_frame
    lines_filtered: List[str] = []
    for line in body_lines:
        s = (line or "").strip() if isinstance(line, str) else str(line).strip()
        if s:
            lines_filtered.append(s)
    if not lines_filtered:
        tf.text = "(No content for this section.)"
    else:
        tf.text = _truncate(lines_filtered[0], 2000)
        for line in lines_filtered[1:]:
            p = tf.add_paragraph()
            p.text = _truncate(line, 2000)
            p.level = 0
            p.font.size = Pt(12)


def _bullets_slide(prs: Presentation, title: str, bullets: Iterable[str]) -> None:
    lines: List[str] = []
    for b in bullets:
        t = (b or "").strip()
        if t:
            lines.append(f"• {_truncate(t, 280)}")
    if not lines:
        lines = ["• (none)"]
    _add_title_content_slide(prs, title, lines[:40])


def build_ppt_bytes(report: AnalysisReport) -> tuple[bytes, str]:
    """
    Returns (pptx_bytes, suggested_filename).
    """
    prs = Presentation()

    # Title slide (layout 0)
    slide0 = prs.slides.add_slide(prs.slide_layouts[0])
    slide0.shapes.title.text = _truncate(report.title, 200)
    if len(slide0.placeholders) > 1:
        sub = slide0.placeholders[1]
        sub.text = _truncate(
            report.description or "Project Intel — Intelligence report",
            600,
        )

    # Similarity overview
    sim_lines = [
        f"Similarity score: {report.similarity_score}%",
        f"Label: {report.similarity_label or '—'}",
        f"HF live scoring: {report.similarity_hf_live if report.similarity_hf_live is not None else '—'}",
    ]
    if report.similarity_description:
        sim_lines.append("")
        sim_lines.append(_truncate(report.similarity_description, 900))
    _add_title_content_slide(prs, "Similarity overview", sim_lines)

    # SWOT
    swot_lines: List[str] = []
    for label, items in (
        ("Strengths", report.swot.strengths),
        ("Weaknesses", report.swot.weaknesses),
        ("Opportunities", report.swot.opportunities),
        ("Threats", report.swot.threats),
    ):
        swot_lines.append(f"{label}:")
        if items:
            for it in items:
                swot_lines.append(f"  • {_truncate(it, 220)}")
        else:
            swot_lines.append("  • (none)")
        swot_lines.append("")
    _add_title_content_slide(prs, "SWOT", swot_lines)

    # Competitors
    if report.found_projects:
        comp_lines: List[str] = []
        for i, p in enumerate(report.found_projects[:12], 1):
            comp_lines.append(f"{i}. {_truncate(p.name, 120)}")
            if p.snippet:
                comp_lines.append(f"   {_truncate(p.snippet, 240)}")
            if p.link:
                comp_lines.append(f"   {_truncate(p.link, 200)}")
            comp_lines.append("")
        _add_title_content_slide(prs, "Competitors found", comp_lines)
    else:
        _add_title_content_slide(prs, "Competitors found", ["No competitor rows in this report."])

    # Tech comparison
    if report.tech_comparison:
        tech_lines: List[str] = []
        for row in report.tech_comparison[:20]:
            tech_lines.append(f"• {_truncate(row.area or 'Area', 80)}")
            tech_lines.append(f"  Yours: {_truncate(row.user_stack, 200)}")
            tech_lines.append(f"  Market: {_truncate(row.competitor_stack, 200)}")
            tech_lines.append(f"  Note: {_truncate(row.advantage, 220)}")
            tech_lines.append("")
        _add_title_content_slide(prs, "Tech lens", tech_lines)
    else:
        _add_title_content_slide(prs, "Tech lens", ["No tech comparison rows for this run."])

    # Recommendations
    _bullets_slide(prs, "Recommendations", report.recommendations or ["(none)"])

    # Strategy
    strat_lines: List[str] = ["Positioning & strategy", ""]
    v = report.strategy.venn
    strat_lines.append("Shared features:")
    for x in (v.shared_features or [])[:12]:
        strat_lines.append(f"  • {_truncate(x, 200)}")
    strat_lines.append("")
    strat_lines.append("Unique to you:")
    for x in (v.unique_to_you or [])[:12]:
        strat_lines.append(f"  • {_truncate(x, 200)}")
    strat_lines.append("")
    strat_lines.append("Unique to market:")
    for x in (v.unique_to_market or [])[:12]:
        strat_lines.append(f"  • {_truncate(x, 200)}")
    strat_lines.append("")
    strat_lines.append("Market pivots:")
    for x in (report.strategy.market_pivots or [])[:10]:
        strat_lines.append(f"  • {_truncate(x, 220)}")
    strat_lines.append("")
    strat_lines.append("Monetization:")
    for x in (report.strategy.monetization or [])[:10]:
        strat_lines.append(f"  • {_truncate(x, 220)}")
    _add_title_content_slide(prs, "Strategy", strat_lines)

    # Devil's advocate
    _bullets_slide(prs, "Devil's advocate", report.devils_advocate or ["(none)"])

    # References
    if report.references:
        ref_lines: List[str] = []
        for r in report.references[:25]:
            ref_lines.append(f"• {_truncate(r.title, 160)} ({r.kind})")
            ref_lines.append(f"  {_truncate(r.url, 240)}")
            ref_lines.append("")
        _add_title_content_slide(prs, "References", ref_lines)
    else:
        _add_title_content_slide(prs, "References", ["(none)"])

    # Market snippet + excerpt
    tail: List[str] = []
    if report.market_search_snippet:
        tail.append("Market / web context (excerpt):")
        tail.append(_truncate(report.market_search_snippet, 2500))
        tail.append("")
    tail.append("Document excerpt (merged):")
    tail.append(_truncate(report.file_content, 3500))
    _add_title_content_slide(prs, "Context excerpts", tail)

    bio = BytesIO()
    prs.save(bio)
    bio.seek(0)
    return bio.getvalue(), _safe_filename(report.title)
