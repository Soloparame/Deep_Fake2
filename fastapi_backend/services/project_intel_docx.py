"""Build a Word document from a Project Intel AnalysisReport using python-docx."""

from __future__ import annotations

import re
from io import BytesIO
from typing import Iterable, List

from docx import Document
from docx.shared import Pt

from fastapi_backend.schemas.analysis import AnalysisReport


def _truncate(text: str, max_len: int = 500) -> str:
    text = (text or "").strip()
    if len(text) <= max_len:
        return text
    return text[: max_len - 1] + "…"


def _safe_filename(title: str) -> str:
    base = re.sub(r"[^\w\-]+", "_", (title or "report").strip())[:60] or "report"
    return f"{base}.docx"


def _add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(_truncate(text, 200), level=level)


def _add_bullets(doc: Document, items: Iterable[str]) -> None:
    any_item = False
    for raw in items:
        t = (raw or "").strip()
        if not t:
            continue
        any_item = True
        p = doc.add_paragraph(_truncate(t, 2000), style="List Bullet")
        for run in p.runs:
            run.font.size = Pt(11)
    if not any_item:
        p = doc.add_paragraph("(none)", style="List Bullet")
        for run in p.runs:
            run.font.size = Pt(11)


def _add_body_lines(doc: Document, lines: List[str]) -> None:
    filtered = [((line or "").strip() if isinstance(line, str) else str(line).strip()) for line in lines]
    filtered = [x for x in filtered if x]
    if not filtered:
        doc.add_paragraph("(No content for this section.)")
        return
    for line in filtered:
        p = doc.add_paragraph(_truncate(line, 4000))
        for run in p.runs:
            run.font.size = Pt(11)


def build_docx_bytes(report: AnalysisReport) -> tuple[bytes, str]:
    doc = Document()

    doc.add_heading(_truncate(report.title, 200), 0)
    sub = doc.add_paragraph(_truncate(report.description or "Project Intel — Intelligence report", 1200))
    for run in sub.runs:
        run.font.size = Pt(12)

    _add_heading(doc, "Similarity overview", 1)
    sim_lines = [
        f"Similarity score: {report.similarity_score}%",
        f"Label: {report.similarity_label or '—'}",
        f"HF live scoring: {report.similarity_hf_live if report.similarity_hf_live is not None else '—'}",
    ]
    if report.similarity_description:
        sim_lines.append(_truncate(report.similarity_description, 2000))
    _add_body_lines(doc, sim_lines)

    _add_heading(doc, "SWOT", 1)
    for label, items in (
        ("Strengths", report.swot.strengths),
        ("Weaknesses", report.swot.weaknesses),
        ("Opportunities", report.swot.opportunities),
        ("Threats", report.swot.threats),
    ):
        doc.add_paragraph(f"{label}:", style="Heading 3")
        _add_bullets(doc, items if items else ["(none)"])

    _add_heading(doc, "Competitors found", 1)
    if report.found_projects:
        comp_lines: List[str] = []
        for i, p in enumerate(report.found_projects[:20], 1):
            comp_lines.append(f"{i}. {_truncate(p.name, 200)}")
            if p.snippet:
                comp_lines.append(f"   {_truncate(p.snippet, 500)}")
            if p.link:
                comp_lines.append(f"   {_truncate(p.link, 300)}")
            comp_lines.append("")
        _add_body_lines(doc, comp_lines)
    else:
        doc.add_paragraph("No competitor rows in this report.")

    _add_heading(doc, "Tech lens", 1)
    if report.tech_comparison:
        tech_lines: List[str] = []
        for row in report.tech_comparison[:25]:
            tech_lines.append(f"{_truncate(row.area or 'Area', 120)}")
            tech_lines.append(f"  Yours: {_truncate(row.user_stack, 400)}")
            tech_lines.append(f"  Market: {_truncate(row.competitor_stack, 400)}")
            tech_lines.append(f"  Note: {_truncate(row.advantage, 400)}")
            tech_lines.append("")
        _add_body_lines(doc, tech_lines)
    else:
        doc.add_paragraph("No tech comparison rows for this run.")

    _add_heading(doc, "Recommendations", 1)
    _add_bullets(doc, report.recommendations or ["(none)"])

    _add_heading(doc, "Strategy", 1)
    v = report.strategy.venn
    doc.add_paragraph("Shared features:", style="Heading 3")
    _add_bullets(doc, v.shared_features or ["(none)"])
    doc.add_paragraph("Unique to you:", style="Heading 3")
    _add_bullets(doc, v.unique_to_you or ["(none)"])
    doc.add_paragraph("Unique to market:", style="Heading 3")
    _add_bullets(doc, v.unique_to_market or ["(none)"])
    doc.add_paragraph("Market pivots:", style="Heading 3")
    _add_bullets(doc, report.strategy.market_pivots or ["(none)"])
    doc.add_paragraph("Monetization:", style="Heading 3")
    _add_bullets(doc, report.strategy.monetization or ["(none)"])

    _add_heading(doc, "Devil's advocate", 1)
    _add_bullets(doc, report.devils_advocate or ["(none)"])

    _add_heading(doc, "References", 1)
    if report.references:
        for r in report.references[:30]:
            line = f"{r.title} ({r.kind}) — {r.url}"
            doc.add_paragraph(_truncate(line, 600), style="List Bullet")
    else:
        doc.add_paragraph("(none)")

    _add_heading(doc, "Context excerpts", 1)
    if report.market_search_snippet:
        doc.add_paragraph("Market / web context (excerpt):", style="Heading 3")
        doc.add_paragraph(_truncate(report.market_search_snippet, 8000))
    doc.add_paragraph("Document excerpt (merged):", style="Heading 3")
    doc.add_paragraph(_truncate(report.file_content, 12000))

    bio = BytesIO()
    doc.save(bio)
    bio.seek(0)
    return bio.getvalue(), _safe_filename(report.title)
